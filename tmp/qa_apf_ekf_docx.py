import json
import re
import sys
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "paper" / "COLREG_Compliant_ASV_Navigation_Yida_Zhu_Final_Cited_Flowchart_APF_EKF.docx"
REPORT = ROOT / "paper" / "qa_apf_ekf.json"

doc = Document(DOCX)
paragraph_text = [paragraph.text.strip() for paragraph in doc.paragraphs]
full_text = "\n".join(paragraph_text)
captions = [
    paragraph.text.strip()
    for paragraph in doc.paragraphs
    if paragraph.style.name == "Caption" and paragraph.text.strip().startswith("Figure ")
]
figure_numbers = [
    int(match.group(1))
    for caption in captions
    if (match := re.match(r"Figure\s+(\d+)\.", caption))
]
equation_numbers = sorted(
    {
        int(number)
        for number in re.findall(r"\(3\.(\d+)\)", full_text)
        if 1 <= int(number) <= 14
    }
)
required_phrases = [
    "constant-turn-rate-and-velocity (CTRV)",
    "Pₖ⁺ = (I − KₖH)Pₖ⁻(I − KₖH)ᵀ + KₖRₖKₖᵀ",
    "λⱼ(q) =",
    "F_res = F_goal + F_path + ΣⱼF_rep,ⱼ + F_rule",
    "Author-generated schematic of the size-aware predictive artificial potential field",
    "(Kalman, 1960; Xie et al., 2024)",
    "(Lyu et al., 2024; Li et al., 2024a)",
]
required_references = [
    "Khatib, O. (1986)",
    "Kalman, R.E. (1960)",
    "Lyu, H., Liu, W.",
    "Xie, Y., Nanlal, C. and Liu, Y. (2024)",
]

report = {
    "docx": str(DOCX),
    "paragraph_count": len(doc.paragraphs),
    "table_count": len(doc.tables),
    "inline_shape_count": len(doc.inline_shapes),
    "figure_caption_count": len(captions),
    "figure_numbers": figure_numbers,
    "figure_sequence_ok": figure_numbers == list(range(1, 20)),
    "equation_numbers": equation_numbers,
    "equation_sequence_ok": equation_numbers == list(range(1, 15)),
    "required_phrases": {phrase: phrase in full_text for phrase in required_phrases},
    "required_references": {reference: reference in full_text for reference in required_references},
    "unresolved_markers": [text for text in paragraph_text if "[[" in text or "]]" in text],
}
report["passed"] = (
    report["figure_sequence_ok"]
    and report["equation_sequence_ok"]
    and all(report["required_phrases"].values())
    and all(report["required_references"].values())
    and not report["unresolved_markers"]
    and report["inline_shape_count"] == 19
    and report["table_count"] == 1
)

REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
sys.stdout.reconfigure(encoding="utf-8")
print(json.dumps(report, ensure_ascii=False, indent=2))
