from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import build_thesis_docx as base


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "paper" / "colreg_asv_thesis_en_cited.md"
OUTPUT = ROOT / "paper" / "COLREG_Compliant_ASV_Navigation_Yida_Zhu_Final_Cited_Flowchart_APF_EKF.docx"

PSEUDOCODE = [
    ("Input", "tracked LiDAR obstacles O, predicted virtual obstacles V, own-vessel velocity v_own, time t, tracking command u_track"),
    ("1", "v_own <- current_velocity_body()"),
    ("2", "candidates <- lidar_obstacles union update_apf_virtual_obstacles()"),
    ("3", "best <- {rule: none, controller: default_apf, score: (priority(none), infinity, infinity)}"),
    ("4", "for each obstacle o in candidates do"),
    ("5", "    p <- obstacle position in the own-vessel body frame; continue if p is invalid"),
    ("6", "    if o is observed and outside the forward activation sector then continue"),
    ("7", "    v <- obstacle body-frame velocity from o or from its stable track; otherwise use zero"),
    ("8", "    rule <- encounter_mode(o) for a virtual obstacle, otherwise classify_encounter(p, v, v_own)"),
    ("9", "    if rule is not head_on, overtaking, or crossing then rule <- static_obstacle"),
    ("10", "    (TCPA, DCPA) <- apf_cpa_metrics(p, v, v_own)"),
    ("11", "    score <- (priority(rule), finite(TCPA) ? TCPA : infinity, finite(DCPA) ? DCPA : norm(p))"),
    ("12", "    if score < best.score then best <- {rule, TCPA, DCPA, score}"),
    ("13", "selected_rule <- best.rule"),
    ("14", "if selected_rule is head_on or overtaking then branch <- overtaking_head_on"),
    ("15", "else if selected_rule is crossing then branch <- crossing"),
    ("16", "else branch <- default_apf"),
    ("17", "u_cmd <- branch.compute_apf_control(t, u_track)"),
    ("18", "record branch, encounter mode, TCPA, and DCPA in controller state and snapshots"),
    ("19", "return u_cmd"),
]


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(Cm(width_cm).emu / 635)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "bottom", "insideH"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6" if edge != "insideH" else "2")
        tag.set(qn("w:color"), "000000" if edge != "insideH" else "BFBFBF")
    for edge in ("left", "right", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")


def add_algorithm_before(doc: Document, anchor_text: str):
    anchor = next(p for p in doc.paragraphs if p.text.strip() == anchor_text and p.style.name == "Heading 2")

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption.paragraph_format.space_before = Pt(8)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True
    run = caption.add_run("Algorithm 1. COLREG encounter selection and APF controller dispatch")
    base.set_font(run, "Times New Roman", "Times New Roman", 10, bold=True, color=RGBColor(0, 0, 0))

    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    set_table_borders(table)
    for line_no, statement in PSEUDOCODE:
        cells = table.add_row().cells
        set_cell_width(cells[0], 1.1)
        set_cell_width(cells[1], 14.3)
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        for cell in cells:
            base.set_cell_margins(cell, top=55, start=90, bottom=55, end=90)
        p0, p1 = cells[0].paragraphs[0], cells[1].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p0.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r0, r1 = p0.add_run(line_no), p1.add_run(statement)
        base.set_font(r0, "Consolas", "Consolas", 8.5, bold=(line_no == "Input"), color=RGBColor(0, 0, 0))
        base.set_font(r1, "Consolas", "Consolas", 8.5, bold=(line_no == "Input"), color=RGBColor(0, 0, 0))
    table.rows[0].cells[0]._tc.get_or_add_tcPr().append(OxmlElement("w:shd"))
    table.rows[0].cells[0]._tc.tcPr[-1].set(qn("w:fill"), "E7E6E6")
    table.rows[0].cells[1]._tc.get_or_add_tcPr().append(OxmlElement("w:shd"))
    table.rows[0].cells[1]._tc.tcPr[-1].set(qn("w:fill"), "E7E6E6")
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

    source = doc.add_paragraph(style="Source Note")
    source.paragraph_format.space_before = Pt(3)
    source.paragraph_format.space_after = Pt(8)
    source.paragraph_format.keep_with_next = True
    sr = source.add_run("Source: reconstructed directly from select_colreg_strategy() and compute_apf_control() in src/laptop.py.")
    base.set_font(sr, "Times New Roman", "Times New Roman", 9, italic=True, color=RGBColor(0, 0, 0))

    anchor._p.addprevious(caption._p)
    anchor._p.addprevious(table._tbl)
    anchor._p.addprevious(source._p)


def format_references(doc: Document):
    in_references = False
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "References" and paragraph.style.name == "Heading 1":
            in_references = True
            continue
        if in_references and paragraph.style.name == "Heading 1":
            break
        if in_references and paragraph.text.lstrip().startswith("["):
            paragraph.paragraph_format.first_line_indent = Cm(-0.74)
            paragraph.paragraph_format.left_indent = Cm(0.74)
            paragraph.paragraph_format.space_after = Pt(4)


def build():
    base.SOURCE = SOURCE
    base.OUTPUT = OUTPUT
    base.NAVY = RGBColor(0, 0, 0)
    base.GRAY = RGBColor(0, 0, 0)
    base.build()
    doc = Document(OUTPUT)
    add_algorithm_before(doc, "3.8 Size-aware APF design")
    format_references(doc)
    doc.core_properties.comments = (
        "Citations verified and expanded; APF/EKF principles, equations, and a code-grounded potential-field schematic added on 29 June 2026."
    )
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
