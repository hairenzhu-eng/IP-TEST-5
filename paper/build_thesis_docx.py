from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "paper" / "colreg_asv_thesis_en.md"
OUTPUT = ROOT / "paper" / "COLREG_Compliant_ASV_Navigation_Yida_Zhu_Final.docx"

NAVY = RGBColor(31, 78, 121)
GRAY = RGBColor(90, 90, 90)
TITLE = "COLREG-Compliant Autonomous Surface Vessel Navigation in Busy Waterways"
SUBTITLE = "A Webots Study Based on LiDAR Clustering, Extended Kalman Prediction, and Size-Aware Artificial Potential Fields"
SHORT_TITLE = "COLREG-Compliant ASV Navigation"


def set_font(run, latin="Times New Roman", east_asia="Times New Roman", size=None, bold=None, italic=None, color=None):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_style_font(style, latin, east_asia, size, bold=False, italic=False, color=None):
    style.font.name = latin
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    if color:
        style.font.color.rgb = color


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.append(fld)
    set_font(run, size=9, color=GRAY)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    set_style_font(normal, "Times New Roman", "Times New Roman", 10.5)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True

    h1 = doc.styles["Heading 1"]
    set_style_font(h1, "Arial", "Arial", 16, bold=True, color=NAVY)
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    set_style_font(h2, "Arial", "Arial", 13, bold=True, color=NAVY)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    set_style_font(caption, "Times New Roman", "Times New Roman", 9.5, color=GRAY)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.keep_with_next = False

    if "Source Note" not in [s.name for s in doc.styles]:
        note = doc.styles.add_style("Source Note", WD_STYLE_TYPE.PARAGRAPH)
        set_style_font(note, "Times New Roman", "Times New Roman", 9.5, italic=True, color=GRAY)
        note.paragraph_format.left_indent = Cm(0.7)
        note.paragraph_format.right_indent = Cm(0.7)
        note.paragraph_format.space_before = Pt(6)
        note.paragraph_format.space_after = Pt(8)
        note.paragraph_format.line_spacing = 1.25

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = hp.add_run("FEEG6012 MSc Project  |  COLREG ASV Navigation")
    set_font(hr, "Arial", "Arial", 8.5, color=GRAY)
    add_page_number(section.footer.paragraphs[0])


def add_cover(doc: Document):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("UNIVERSITY OF SOUTHAMPTON")
    set_font(r, "Arial", "Arial", 12, bold=True, color=GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(TITLE)
    set_font(r, "Arial", "Arial", 20, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run(SUBTITLE)
    set_font(r, "Arial", "Arial", 12.5, color=GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run(
        "A dissertation submitted to the Faculty of Engineering and Physical Sciences\n"
        "in partial fulfilment of the requirements for the degree of Master of Science"
    )
    set_font(r, size=11)

    for label, value in (
        ("Author", "Yida Zhu"),
        ("Student ID", "33393559"),
        ("Module", "FEEG6012 MSc Project"),
        ("Programme", "MSc project dissertation"),
        ("Date", "June 2026"),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(5)
        a = p.add_run(f"{label}：")
        set_font(a, "Arial", "Arial", 11, bold=True)
        b = p.add_run(value)
        set_font(b, size=11)

    doc.add_paragraph()
    p = doc.add_paragraph(style="Source Note")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "Research dissertation draft developed from the implemented controller code, Webots experiment logs, the approved scoping study, and departmental thesis examples. Final submission should be checked by the author and accompanied by the required GenAI declaration."
    )
    doc.add_page_break()


def add_declaration_page(doc: Document):
    p = doc.add_paragraph("Declaration", style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    intro = (
        "I, Yida Zhu, declare that this dissertation and the work presented in it are my own. "
        "This document has been prepared from my MSc project materials, controller implementation, "
        "simulation outputs, and subsequent analysis."
    )
    add_body_paragraph(doc, intro)

    declarations = [
        "This work was completed during my candidature for the MSc project.",
        "Where the ideas, data, methods, or words of other authors have been used, they have been acknowledged through citation.",
        "The dissertation adapts material from my approved scoping study as project background and development context, but the present analysis, structure, and conclusions have been rewritten for the final thesis.",
        "Claims about implemented functionality are limited to modules that are supported by code inspection, logs, figures, and reproducible post-processing artefacts.",
        "Any use of generative AI in drafting, editing, or document preparation must be declared by the author in the final university submission paperwork.",
    ]
    for idx, item in enumerate(declarations, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.3)
        p.paragraph_format.first_line_indent = Cm(-0.3)
        run = p.add_run(f"{idx}. {item}")
        set_font(run, size=10.5)
    doc.add_page_break()


def add_acknowledgements_page(doc: Document):
    p = doc.add_paragraph("Acknowledgements", style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    text = (
        "The author acknowledges the supervisory guidance, the project specification provided through the "
        "scoping stage, and the Webots-based experimentation environment that made the present analysis possible. "
        "The author also acknowledges that the dissertation benefited from prior project artefacts, including "
        "simulation logs, plotting scripts, and departmental thesis examples used only as formatting references."
    )
    add_body_paragraph(doc, text)
    doc.add_page_break()


def clean_inline(text: str) -> str:
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    text = text.replace("`", "")
    return text


def add_body_paragraph(doc: Document, text: str, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.first_line_indent = Cm(0.74) if style is None else None
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(clean_inline(text))
    return p


def add_figure(doc: Document, marker: str):
    payload = marker[len("[[FIGURE:") : -2]
    path_text, caption = payload.split("|", 1)
    path = Path(path_text)
    if not path.exists():
        raise FileNotFoundError(path)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()

    from PIL import Image

    with Image.open(path) as image:
        width_px, height_px = image.size
    ratio = height_px / width_px
    max_width = 6.25
    max_height = 7.25
    width = min(max_width, max_height / ratio)
    picture = run.add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("descr", caption.strip())
    picture._inline.docPr.set("title", caption.split(".", 1)[0].strip())

    cap = doc.add_paragraph(style="Caption")
    cap.add_run(caption)


def add_equation(doc: Document, marker: str):
    payload = marker[len("[[EQUATION:") : -2]
    expression, number = payload.split("|", 1)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(7.95), WD_TAB_ALIGNMENT.CENTER)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(15.7), WD_TAB_ALIGNMENT.RIGHT)
    paragraph.add_run("\t")
    equation_run = paragraph.add_run(expression.strip())
    set_font(equation_run, "Cambria Math", "Cambria Math", 10.5)
    paragraph.add_run("\t")
    number_run = paragraph.add_run(number.strip())
    set_font(number_run, "Times New Roman", "Times New Roman", 10.5)


def build():
    text = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_declaration_page(doc)
    add_acknowledgements_page(doc)

    lines = text.splitlines()
    body_start = next(i for i, line in enumerate(lines) if line.strip() == "# Abstract")
    toc_entries = []
    figure_entries = []
    for line in lines[body_start:]:
        if line.startswith("# "):
            title = line[2:].strip()
            if title != "Abstract":
                toc_entries.append((1, title))
        elif line.startswith("## "):
            toc_entries.append((2, line[3:].strip()))
        elif line.startswith("[[FIGURE:") and line.endswith("]]"):
            payload = line[len("[[FIGURE:") : -2]
            _, caption = payload.split("|", 1)
            figure_entries.append(caption.strip())
    pending = []
    last_was_pagebreak = False

    def flush():
        nonlocal pending
        if pending:
            paragraph = " ".join(x.strip() for x in pending if x.strip())
            if paragraph:
                add_body_paragraph(doc, paragraph)
            pending = []

    for raw in lines[body_start:]:
        line = raw.strip()
        if not line:
            flush()
            continue
        if line == "\\pagebreak":
            flush()
            doc.add_page_break()
            last_was_pagebreak = True
            continue
        if line == "[[TOC]]":
            flush()
            doc.add_paragraph("Contents", style="Heading 1")
            for level, title in toc_entries:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.0 if level == 1 else 0.75)
                p.paragraph_format.space_after = Pt(3)
                r = p.add_run(title)
                set_font(r, "Arial", "宋体", 10.5, bold=(level == 1))
            if figure_entries:
                doc.add_paragraph("List of Figures", style="Heading 2")
                for title in figure_entries:
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Cm(0.0)
                    p.paragraph_format.space_after = Pt(3)
                    r = p.add_run(title)
                    set_font(r, size=10)
            last_was_pagebreak = False
            continue
        if line.startswith("[[FIGURE:") and line.endswith("]]" ):
            flush()
            add_figure(doc, line)
            last_was_pagebreak = False
            continue
        if line.startswith("[[EQUATION:") and line.endswith("]]" ):
            flush()
            add_equation(doc, line)
            last_was_pagebreak = False
            continue
        if line.startswith("# "):
            flush()
            title = line[2:].strip()
            if title != "Abstract" and not last_was_pagebreak:
                doc.add_page_break()
            p = doc.add_paragraph(title, style="Heading 1")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            last_was_pagebreak = False
            continue
        if line.startswith("## "):
            flush()
            doc.add_paragraph(line[3:].strip(), style="Heading 2")
            last_was_pagebreak = False
            continue
        if line.startswith("> "):
            flush()
            add_body_paragraph(doc, line[2:], style="Source Note")
            last_was_pagebreak = False
            continue
        pending.append(line)
        last_was_pagebreak = False
    flush()

    props = doc.core_properties
    props.title = TITLE
    props.subject = "FEEG6012 MSc Project dissertation draft"
    props.author = "Yida Zhu"
    props.keywords = "COLREGs, ASV, LiDAR, EKF, APF, Webots"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
